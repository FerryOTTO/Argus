"""Per-format text extraction for media attachments."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any


def _read_text(path: str | Path, encodings: tuple[str, ...]) -> str:
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            return Path(path).read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError, OSError) as exc:
            last_error = exc
    raise ValueError(f"unable to decode file as text: {last_error}")


def extract_pdf(path: str | Path) -> str:
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24
    except ImportError:  # pragma: no cover - older package naming
        import fitz  # type: ignore[no-redef]
    parts: list[str] = []
    with fitz.open(str(path)) as document:
        for page in document:
            text = page.get_text().strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


def extract_docx(path: str | Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def extract_xlsx(path: str | Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [
                    str(cell).strip()
                    for cell in row
                    if cell is not None and str(cell).strip()
                ]
                if cells:
                    parts.append("\t".join(cells))
    finally:
        workbook.close()
    return "\n".join(parts)


def extract_csv(path: str | Path) -> str:
    text = _read_text(path, ("utf-8-sig", "gbk", "utf-8"))
    rows: list[str] = []
    try:
        for row in csv.reader(io.StringIO(text)):
            cells = [cell.strip() for cell in row if cell.strip()]
            if cells:
                rows.append(",".join(cells))
    except (csv.Error, ValueError):
        return text
    return "\n".join(rows) if rows else text


def extract_text(path: str | Path) -> str:
    return _read_text(path, ("utf-8", "utf-8-sig", "gbk", "utf-16", "latin-1"))


def extract_image(path: str | Path) -> str:
    from . import ocr

    return ocr.ocr_image(str(path))


def extract_from_file(
    path: str | Path,
    media_type: str,
    cfg: dict[str, Any],
) -> str:
    """Extract text from a local file, truncated to cfg['max_text_chars']."""
    target = Path(path)
    if not target.exists():
        raise ValueError("attachment file not found")
    if target.stat().st_size > int(cfg["max_file_bytes"]):
        raise ValueError(f"attachment exceeds {cfg['max_file_bytes']} bytes")

    if media_type == "pdf":
        text = extract_pdf(target)
    elif media_type == "docx":
        text = extract_docx(target)
    elif media_type == "xlsx":
        text = extract_xlsx(target)
    elif media_type == "csv":
        text = extract_csv(target)
    elif media_type == "txt":
        text = extract_text(target)
    elif media_type in {"png", "jpeg", "gif", "bmp", "webp"}:
        text = extract_image(target)
    else:
        raise ValueError(f"unsupported media type: {media_type}")

    max_chars = int(cfg["max_text_chars"])
    return text[:max_chars] if len(text) > max_chars else text
